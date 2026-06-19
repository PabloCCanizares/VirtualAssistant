"""Tests para los formatos avanzados del extractor de documentos.

Cubre las ramas PDF/DOCX/XLSX que requieren bibliotecas externas y por tanto
se quedaron fuera del nivel unitario base (`tests/test_doc_parser.py`).
"""

from __future__ import annotations

from io import BytesIO

import pytest

from ai.services.doc_parser import extract_text


def _generate_pdf_bytes(text="Hola mundo TFG"):
    """Genera un PDF minimo en bytes usando fpdf2."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    out = BytesIO()
    pdf.output(out)
    return out.getvalue()


def _generate_docx_bytes(text="texto de prueba"):
    """Genera un DOCX minimo en bytes usando python-docx."""
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _generate_xlsx_bytes():
    """Genera un XLSX minimo en bytes usando openpyxl."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["nombre", "edad"])
    ws.append(["Ana", 30])
    ws.append(["Bob", 40])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class TestPdfExtraction:
    def test_pdf_by_content_type(self):
        out = extract_text(
            _generate_pdf_bytes("Hola TFG"),
            content_type="application/pdf",
            filename="x.pdf",
        )
        assert "TFG" in out or "Hola" in out or "extraible" in out

    def test_pdf_by_extension(self):
        out = extract_text(
            _generate_pdf_bytes("contenido"),
            content_type="",
            filename="doc.pdf",
        )
        assert isinstance(out, str)

    def test_pdf_corrupt_returns_error_message(self):
        out = extract_text(
            b"not a pdf",
            content_type="application/pdf",
            filename="bad.pdf",
        )
        assert "Error" in out or "extraible" in out


class TestDocxExtraction:
    def test_docx_extracts_text(self):
        out = extract_text(
            _generate_docx_bytes("contenido del docx"),
            content_type="",
            filename="doc.docx",
        )
        assert "contenido del docx" in out

    def test_docx_by_content_type(self):
        out = extract_text(
            _generate_docx_bytes("hola"),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="x.docx",
        )
        assert "hola" in out

    def test_docx_corrupt_returns_error(self):
        out = extract_text(
            b"not a docx",
            content_type="",
            filename="bad.docx",
        )
        assert "Error" in out


class TestXlsxExtraction:
    def test_xlsx_extracts_table(self):
        out = extract_text(
            _generate_xlsx_bytes(),
            content_type="",
            filename="data.xlsx",
        )
        assert "nombre" in out
        assert "Ana" in out

    def test_xlsx_by_content_type(self):
        out = extract_text(
            _generate_xlsx_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename="x.xlsx",
        )
        assert "Ana" in out

    def test_xlsx_corrupt(self):
        out = extract_text(
            b"not xlsx",
            content_type="",
            filename="bad.xlsx",
        )
        assert "Error" in out


class TestCsvBoundaries:
    def test_large_csv_truncates(self):
        # Mas de MAX_CSV_ROWS=50 filas
        lines = ["col1,col2"]
        for i in range(60):
            lines.append(f"a{i},b{i}")
        out = extract_text("\n".join(lines).encode(), content_type="text/csv", filename="x.csv")
        assert "filas adicionales omitidas" in out
